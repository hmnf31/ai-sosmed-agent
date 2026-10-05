"""Test pembacaan spreadsheet publik Google Sheets tanpa kredensial.

Sheet publik bisa dibaca lewat endpoint export .xlsx. Karena itu test ini
memalsukan respons HTTP, bukan memanggil jaringan.
"""
import io

import pytest

from utils.chess import spreadsheet

openpyxl = pytest.importorskip("openpyxl")


@pytest.fixture(autouse=True)
def _no_other_sources(monkeypatch):
    monkeypatch.delenv("SHEET_CREDENTIALS_JSON", raising=False)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    monkeypatch.delenv("SHEET_XLSX_PATH", raising=False)
    spreadsheet._PUBLIC_CACHE.clear()


def _workbook_bytes(sheets):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


class _Response:
    def __init__(self, content=b"", status_code=200):
        self.content = content
        self.status_code = status_code


def _serve(monkeypatch, sheets):
    payload = _workbook_bytes(sheets)
    monkeypatch.setenv("SHEET_PUBLIC_ID", "SHEET_ID_ABC")
    monkeypatch.setattr(
        spreadsheet.requests, "get", lambda *a, **k: _Response(payload)
    )


def test_public_id_accepts_plain_id_and_full_url(monkeypatch):
    monkeypatch.setenv("SHEET_PUBLIC_ID", "11Q3AIGofm1ZQ")
    assert spreadsheet._public_id() == "11Q3AIGofm1ZQ"

    monkeypatch.setenv(
        "SHEET_PUBLIC_ID",
        "https://docs.google.com/spreadsheets/d/11Q3AIGofm1ZQ/edit?usp=sharing",
    )
    assert spreadsheet._public_id() == "11Q3AIGofm1ZQ"


def test_public_sheet_counts_as_configured(monkeypatch):
    monkeypatch.setenv("SHEET_PUBLIC_ID", "SHEET_ID_ABC")
    assert spreadsheet.is_configured() is True


def test_standings_tab_is_read_as_league_source(monkeypatch):
    _serve(
        monkeypatch,
        {
            "Standings": [
                ["league", "position", "player_id", "name", "mp", "w", "d", "l", "points"],
                ["Liga 1", 1.0, "p041", "Sultan", 2.0, 2.0, 0.0, 0.0, 4.0],
                ["Liga 1", 2.0, "p012", "Mr.G", 1.0, 1.0, 0.0, 0.0, 2.0],
                ["Liga 2", 1.0, "p005", "Ayyme", 1.0, 1.0, 0.0, 0.0, 2.0],
            ]
        },
    )
    rows, source = spreadsheet.get_league_standings("Liga 1")
    assert source == "google_public:Standings"
    assert [r["nama"] for r in rows] == ["Sultan", "Mr.G"]
    assert rows[0]["poin"] == 4
    assert rows[0]["main"] == 2


def test_public_league_filter_is_case_insensitive(monkeypatch):
    _serve(
        monkeypatch,
        {
            "Standings": [
                ["league", "name", "points"],
                ["Liga 1", "Sultan", 4.0],
            ]
        },
    )
    rows, _ = spreadsheet.get_league_standings("liga 1")
    assert [r["nama"] for r in rows] == ["Sultan"]


def test_player_id_column_does_not_steal_nama(monkeypatch):
    """`player_id` cocok sebagian dengan alias `player`; `name` harus tetap jadi nama."""
    _serve(
        monkeypatch,
        {
            "Standings": [
                ["league", "position", "player_id", "name", "username", "points"],
                ["Liga 1", 1.0, "p041", "Sultan", "SultanAulia", 4.0],
            ]
        },
    )
    rows, _ = spreadsheet.get_league_standings()
    assert rows[0]["nama"] == "Sultan"


def test_single_letter_columns_only_match_exactly(monkeypatch):
    """`wo_count` diawali "w" tapi bukan kolom menang."""
    _serve(
        monkeypatch,
        {
            "Standings": [
                ["league", "name", "w", "d", "l", "wo_count", "points"],
                ["Liga 1", "Roy", 0.0, 0.0, 2.0, 1.0, 0.0],
            ]
        },
    )
    rows, _ = spreadsheet.get_league_standings()
    assert rows[0]["menang"] == 0
    assert rows[0]["kalah"] == 2


def test_decimal_points_are_not_read_as_tens(monkeypatch):
    _serve(
        monkeypatch,
        {"Standings": [["league", "name", "points"], ["Liga 1", "Sultan", 4.0]]},
    )
    rows, _ = spreadsheet.get_league_standings()
    assert rows[0]["poin"] == 4


def test_workbook_downloaded_once_per_ttl(monkeypatch):
    _serve(
        monkeypatch,
        {"Standings": [["league", "name", "points"], ["Liga 1", "Sultan", 4.0]]},
    )
    calls = []

    def fake_get(*args, **kwargs):
        calls.append(1)
        payload = _workbook_bytes(
            {"Standings": [["league", "name", "points"], ["Liga 1", "Sultan", 4.0]]}
        )
        return _Response(payload)

    monkeypatch.setattr(spreadsheet.requests, "get", fake_get)
    spreadsheet.get_league_standings()
    spreadsheet.get_league_standings("Liga 1")
    assert len(calls) == 1


def test_public_sheet_without_needed_tab_falls_through(monkeypatch):
    """Tab TCO tidak ada di sheet publik, jadi range itu lewat ke sumber lain."""
    _serve(monkeypatch, {"Standings": [["league", "name", "points"], ["Liga 1", "A", 4.0]]})
    rows, source = spreadsheet.read_range(spreadsheet.RANGE_TCO)
    assert rows == []
    assert source == "tidak_ada"


def test_html_response_reports_permission_problem(monkeypatch):
    monkeypatch.setenv("SHEET_PUBLIC_ID", "SHEET_ID_ABC")
    monkeypatch.setattr(
        spreadsheet.requests,
        "get",
        lambda *a, **k: _Response(b"<html>Sign in</html>"),
    )
    with pytest.raises(ValueError) as exc:
        spreadsheet._public_tables()
    assert "dibagikan" in str(exc.value)


def test_http_error_reports_status(monkeypatch):
    monkeypatch.setenv("SHEET_PUBLIC_ID", "SHEET_ID_ABC")
    monkeypatch.setattr(
        spreadsheet.requests, "get", lambda *a, **k: _Response(b"", 404)
    )
    with pytest.raises(ValueError) as exc:
        spreadsheet._public_tables()
    assert "404" in str(exc.value)


def test_available_tabs_lists_public_sheet(monkeypatch):
    _serve(
        monkeypatch,
        {
            "Standings": [["league", "name", "points"], ["Liga 1", "A", 4.0]],
            "Schedules": [["id", "league", "round"], ["s1", "Liga 1", 3]],
        },
    )
    tabs = spreadsheet.available_tabs()
    assert tabs["google_public"] == ["schedules", "standings"]


def test_tco_tab_alias_jadwal_is_recognized(monkeypatch):
    _serve(
        monkeypatch,
        {
            "Jadwal": [
                ["Tanggal", "Waktu", "Format", "Link"],
                ["2026-10-14", "20.00 WIB", "Blitz", "https://chess.com/tco"],
            ]
        },
    )
    from datetime import datetime

    schedule = spreadsheet.get_tco_schedule(datetime(2026, 10, 14))
    assert schedule["found"] is True
    assert schedule["source"] == "google_public:Jadwal"


def test_local_xlsx_used_when_public_lacks_tco_tab(tmp_path, monkeypatch):
    """Sumber publik melayani liga, file lokal melayani TCO. Keduanya menyatu."""
    _serve(monkeypatch, {"Standings": [["league", "name", "points"], ["Liga 1", "A", 4.0]]})

    local = tmp_path / "club.xlsx"
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheet = workbook.create_sheet("TCO")
    sheet.append(["Tanggal", "Waktu", "Link"])
    sheet.append(["2026-10-14", "20.00 WIB", "https://chess.com/tco"])
    workbook.save(str(local))
    monkeypatch.setenv("SHEET_XLSX_PATH", str(local))

    from datetime import datetime

    schedule = spreadsheet.get_tco_schedule(datetime(2026, 10, 14))
    assert schedule["found"] is True
    assert schedule["source"] == "xlsx"

    standings, source = spreadsheet.get_league_standings("Liga 1")
    assert source == "google_public:Standings"
    assert standings[0]["nama"] == "A"