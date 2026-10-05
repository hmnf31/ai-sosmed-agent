"""Test pembacaan file .xlsx lokal dan batas tabel berbasis baris kosong."""
from datetime import datetime

import pytest

from utils.chess import spreadsheet

openpyxl = pytest.importorskip("openpyxl")


def _write_workbook(path, sheets):
    from openpyxl import Workbook

    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        worksheet = workbook.create_sheet(name)
        for row in rows:
            worksheet.append(row)
    workbook.save(str(path))
    return str(path)


def test_xlsx_is_detected(tmp_path, monkeypatch):
    target = _write_workbook(tmp_path / "club.xlsx", {"TCO": [["Tanggal", "Waktu"]]})
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CREDENTIALS_JSON", raising=False)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    assert spreadsheet.is_configured()


def test_xlsx_tco_schedule(tmp_path, monkeypatch):
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "TCO": [
                ["Tanggal", "Waktu", "Format", "Lokasi", "Link", "Keterangan"],
                ["2026-10-14", "20.00 WIB", "Blitz", "Online", "https://chess.com/tco", ""],
            ]
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    schedule = spreadsheet.get_tco_schedule(datetime(2026, 10, 14))
    assert schedule["found"] is True
    assert schedule["source"] == "xlsx"
    assert schedule["waktu"] == "20.00 WIB"


def test_xlsx_league_standings(tmp_path, monkeypatch):
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "Liga": [
                ["Liga", "Rank", "Nama", "Main", "Menang", "Seri", "Kalah", "Poin"],
                ["A", 2, "Player B", 5, 4, 0, 1, 12],
                ["A", 1, "Player A", 5, 4, 1, 0, 13],
            ]
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    rows, source = spreadsheet.get_league_standings("A")
    assert source == "xlsx"
    assert [r["nama"] for r in rows] == ["Player A", "Player B"]


def test_xlsx_arena_settings(tmp_path, monkeypatch):
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "Arena": [
                ["Link Klub", "https://chess.com/ak"],
                ["Jam", "22.30 WIB"],
                ["Format", "Blitz 5 menit"],
                ["Lokasi", "Online"],
                ["Multiklub", "Ya"],
            ]
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    settings, source = spreadsheet.read_settings()
    assert source == "xlsx"
    assert settings["link_klub"] == "https://chess.com/ak"
    assert settings["jam"] == "22.30 WIB"


def test_xlsx_arena_notes_below_settings_are_ignored(tmp_path, monkeypatch):
    """Baris kosong memisahkan isi dari catatan pada sheet konfigurasi."""
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "Arena": [
                ["Link Klub", "https://chess.com/ak"],
                [],
                ["CARA PAKAI SHEET ARENA", ""],
                ["Link Klub", "https://carangan"],
            ]
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    settings, _ = spreadsheet.read_settings()
    assert settings == {"link_klub": "https://chess.com/ak"}


def test_notes_below_table_are_not_data(tmp_path, monkeypatch):
    """Baris kosong mengakhiri tabel, jadi catatan di bawahnya tidak jadi data."""
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "Liga": [
                ["Liga", "Rank", "Nama", "Poin"],
                ["A", 1, "Player A", 13],
                [],
                ["CARA PAKAI SHEET LIGA", "", "", ""],
                ["Liga ditulis A atau B", "", "", ""],
            ]
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    rows, _ = spreadsheet.get_league_standings()
    assert [r["nama"] for r in rows] == ["Player A"]


def test_readme_sheet_is_skipped_when_sheet_name_missing(tmp_path, monkeypatch):
    target = _write_workbook(
        tmp_path / "club.xlsx",
        {
            "Petunjuk": [["Tujuan", "isi"]],
            "Liga": [["Liga", "Rank", "Nama", "Poin"], ["A", 1, "Player A", 13]],
        },
    )
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    rows, _ = spreadsheet.get_league_standings()
    assert [r["nama"] for r in rows] == ["Player A"]


def test_template_script_creates_expected_sheets(tmp_path):
    from openpyxl import load_workbook

    from scripts.make_sheet_template import build_template

    target = tmp_path / "template.xlsx"
    build_template(str(target))
    workbook = load_workbook(str(target))
    try:
        assert {"TCO", "Arena", "Liga", "Petunjuk"} <= set(workbook.sheetnames)
    finally:
        workbook.close()


def test_template_is_readable_by_spreadsheet_module(tmp_path, monkeypatch):
    """Template harus langsung bisa dibaca modul, tanpa penyesuaian kode."""
    from scripts.make_sheet_template import build_template

    target = str(tmp_path / "template.xlsx")
    build_template(target)
    monkeypatch.setenv("SHEET_XLSX_PATH", target)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)

    rows, source = spreadsheet.get_league_standings("Liga 1")
    assert source == "xlsx"
    assert rows, "data contoh liga harus terbaca"
    assert all(r["liga"] == "Liga 1" for r in rows)

    settings, _ = spreadsheet.read_settings()
    assert settings, "data contoh arena harus terbaca"
    assert all("contoh" not in str(v).lower() for v in settings.values())


def test_xlsx_missing_file_is_not_configured(monkeypatch):
    monkeypatch.setenv("SHEET_XLSX_PATH", "tidak-ada.xlsx")
    monkeypatch.delenv("SHEET_CREDENTIALS_JSON", raising=False)
    monkeypatch.delenv("SHEET_CSV_PATH", raising=False)
    assert not spreadsheet.is_configured()


def test_csv_takes_priority_over_xlsx(tmp_path, monkeypatch):
    """CSV lokal menang kalau keduanya diset, supaya mudah ditimpa saat uji."""
    import csv as csv_mod

    csv_file = tmp_path / "liga.csv"
    with open(csv_file, "w", newline="", encoding="utf-8") as handle:
        csv_mod.writer(handle).writerows(
            [["Liga", "Rank", "Nama", "Poin"], ["A", 1, "Dari CSV", 9]]
        )
    xlsx_file = _write_workbook(
        tmp_path / "club.xlsx",
        {"Liga": [["Liga", "Rank", "Nama", "Poin"], ["A", 1, "Dari XLSX", 9]]},
    )
    monkeypatch.setenv("SHEET_CSV_PATH", str(csv_file))
    monkeypatch.setenv("SHEET_XLSX_PATH", xlsx_file)
    monkeypatch.delenv("SHEET_CREDENTIALS_JSON", raising=False)

    rows, source = spreadsheet.get_league_standings("A")
    assert source == "csv"
    assert rows[0]["nama"] == "Dari CSV"