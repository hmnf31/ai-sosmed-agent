"""Membuat club-data.xlsx: data asli, terpisah dari template.

Berbeda dengan make_sheet_template.py yang menghasilkan template kosong untuk
di-commit, script ini menghasilkan file DATA lokal (tidak di-commit karena
masuk .gitignore) yang sudah berisi jadwal TCO mingguan berikutnya dan
konfigurasi Arena yang disepakati.

Jadwal TCO adalah turnamen reguler setiap Rabu pukul 20.00 WIB, format
Blitz 3+0, secara daring. Tanggal dihitung otomatis dari hari ini; hapus
baris masa depan yang tidak dipakai.

Arena Kings sudah lengkap di kode; file ini hanya menimpa nilai default
bawaan bila panitia punya link/resmi yang berubah.

Jalankan:
    .venv\\Scripts\\python.exe scripts\\make_tco_data.py
lalu pastikan SHEET_XLSX_PATH=club-data.xlsx di .env, dan jalankan
scripts/check_sheet_data.py untuk konfirmasi.
"""
import os
import sys
from datetime import datetime, timedelta, timezone

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

OUTPUT_NAME = "club-data.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
HEADER_FONT = Font(color="FFFFFF", bold=True, size=11)
THIN = Side(style="thin", color="C8D0DA")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

WIB = timezone(timedelta(hours=7))


def _wib_now():
    return datetime.now(WIB)


def _next_wednesdays(count=8, when=None):
    base = (when or _wib_now()).date()
    # Rabu = 2 di Python (Senin=0)
    days_ahead = (2 - base.weekday()) % 7
    if days_ahead == 0:
        days_ahead = 7  # lewati Rabu yang sudah lewat hari ini
    start = base + timedelta(days=days_ahead)
    for _ in range(count):
        yield start
        start += timedelta(weeks=1)


def _add_table(worksheet, columns, rows):
    for index, (title, width, _) in enumerate(columns, start=1):
        cell = worksheet.cell(row=1, column=index, value=title)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = BORDER
        worksheet.column_dimensions[get_column_letter(index)].width = width
    worksheet.row_dimensions[1].height = 24
    worksheet.freeze_panes = "A2"

    for row_index, values in enumerate(rows, start=2):
        for column_index, value in enumerate(values, start=1):
            cell = worksheet.cell(row=row_index, column=column_index, value=value)
            cell.border = BORDER
            if column_index == len(columns):
                cell.alignment = Alignment(vertical="top", wrap_text=True)
    return len(rows) + 2


def _build_tco(workbook):
    columns = [("Judul", 30, False), ("Mode", 20, False), ("Tanggal", 16, True),
               ("Waktu", 14, False), ("Format", 14, False), ("Lokasi", 22, False),
               ("Link", 46, False), ("Keterangan", 34, False)]
    rows = []
    for date in _next_wednesdays():
        rows.append([
            "TCO Mingguan",
            "Internal Mingguan",
            date.isoformat(),
            "20.00 WIB",
            "Blitz 3+0",
            "Online (Chess.com)",
            "",
            "",
        ])
    worksheet = workbook.create_sheet("TCO")
    _add_table(worksheet, columns, rows)


def _build_arena(workbook):
    settings = [
        ("Link Klub", "https://www.chess.com/club/turnamen-tiktok-chess-online-club"),
        ("Link Arena", ""),
        ("Link Form", "https://forms.gle/1YCVyZMPwpWKN8ZX7"),
        ("Jam", "23.00 WIB"),
        ("Format", "3+0"),
        ("Durasi", "120 Menit"),
        ("Standby", "Standby di Room 1 jam sebelum mulai!"),
        ("Batas Form", "22.30 WIB"),
        ("Kontak 1", "1.Bang Wawan @waone0608"),
        ("Kontak 2", "2.Bang Arif @@~Shakaruby"),
        ("Kontak 3", "3.Bang Sul @~S U L"),
        ("Kontak 4", "4. Bang Teddy @Teddy Sapta Prayoga"),
        ("Keterangan", ""),
    ]
    worksheet = workbook.create_sheet("Arena")
    worksheet.column_dimensions["A"].width = 22
    worksheet.column_dimensions["B"].width = 56
    header_a = worksheet.cell(row=1, column=1, value="Pengaturan")
    header_a.fill, header_a.font, header_a.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="left", vertical="center")
    header_b = worksheet.cell(row=1, column=2, value="Nilai")
    header_b.fill, header_b.font, header_b.alignment = HEADER_FILL, HEADER_FONT, Alignment(horizontal="left", vertical="center")
    worksheet.row_dimensions[1].height = 24
    for index, (key, value) in enumerate(settings, start=2):
        key_cell = worksheet.cell(row=index, column=1, value=key)
        value_cell = worksheet.cell(row=index, column=2, value=value)
        key_cell.font, key_cell.border = Font(bold=True), BORDER
        value_cell.border = BORDER
        key_cell.alignment = value_cell.alignment = Alignment(vertical="top")


def build_data(path=OUTPUT_NAME):
    workbook = Workbook()
    workbook.remove(workbook.active)
    _build_tco(workbook)
    _build_arena(workbook)

    readme = workbook.create_sheet("Petunjuk", 0)
    readme.column_dimensions["A"].width, readme.column_dimensions["B"].width = 26, 86
    lines = [
        ("DATA KLUB CATUR", ""),
        ("", ""),
        ("TCO", "Jadwal Turnamen internal mingguan, setiap Rabu 20.00 WIB."),
        ("Arena", "Konfigurasi link/form/kontak Arena Kings. Tanggal dihitung bot."),
        ("", ""),
        ("Catatan", "Hapus baris 'Petunjuk' setelah membaca, atau biarkan — bot mengabaikannya."),
        ("Cek", "Jalankan scripts\\check_sheet_data.py."),
    ]
    for i, (label, value) in enumerate(lines, start=1):
        left = readme.cell(row=i, column=1, value=label)
        left.font = Font(bold=True, size=14 if i == 1 else 11, color="1F3A5F" if i == 1 else None)
        right = readme.cell(row=i, column=2, value=value)
        right.alignment = Alignment(vertical="top", wrap_text=True)
    workbook.active = 0
    workbook.save(path)
    return path


if __name__ == "__main__":
    target = sys.argv[1] if len(sys.argv) > 1 else OUTPUT_NAME
    saved = build_data(target)
    print(f"[DATA] {saved}")
    print("[DATA] sheet: Petunjuk, TCO, Arena")
    print("[DATA] TCO diisi jadwal mingguan hingga 8 minggu ke depan.")
    print("[DATA] Hubungkan lewat SHEET_XLSX_PATH di .env (SHEET_PUBLIC_ID tetap aktif untuk Liga).")
    print("[DATA] club-data.xlsx tidak di-commit; gunakan make_sheet_template.py untuk template yang di-commit.")
